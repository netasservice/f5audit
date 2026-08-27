"""Opt-in post-processing: ping a report's removal candidates over SSH.

This module is the single, documented exception to the REST-only design
(see CLAUDE.md, "The one SSH exception"). It never takes part in
collection or analysis: it opens an already generated xlsx report, reads
the IPs of the Orphan Nodes sheet (by construction the nodes that are
NOT in use — the removal candidates), pings each unique IPv4 once from
the BIG-IP over SSH, and writes the results back into the workbook.

Safety contract:
- The only remote command ever executed is ping. The two command
  templates below are the only command-building sites, the single
  interpolated value is an ``ipaddress``-validated IPv4 literal, and
  ``_send`` refuses at runtime anything that is not a ping command.
- Ping is read-only ICMP: nothing on the device is created, modified,
  or deleted.
- paramiko is an optional dependency (``pip install 'f5audit[ssh]'``),
  imported lazily inside ``SSHCommandRunner.connect`` so the rest of the
  tool works without it.
"""

from __future__ import annotations

import ipaddress
import logging
import re
from collections.abc import Iterable
from dataclasses import dataclass
from typing import Any, Callable

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .client import F5ClientError
from .report import HEADER_FILL, MAX_COLUMN_WIDTH

logger = logging.getLogger("f5audit.pingcheck")

STATUS_HEADER = "Ping (from F5)"
NOTE_HEADER = "Ping note"
IP_HEADER = "IP"
# IPs come only from Orphan Nodes: the sheet is by construction the list
# of every node whose verdict is not IN USE, i.e. the removal candidates.
ADDRESS_SOURCE_SHEET = "Orphan Nodes"
ENRICHED_SHEETS = ("Inventory", "Orphan Nodes")

STATUS_YES = "YES"
STATUS_NO = "NO"
STATUS_UNKNOWN = "UNKNOWN"
STATUS_NOT_TESTED = "NOT TESTED"

# The only two remote commands this module can ever produce. The bash
# form works for Advanced-shell accounts; tmsh-shell accounts need the
# tmsh form (auto-detected with one probe).
BASH_PING_TEMPLATE = "ping -c {count} -W 1 {ip}"
TMSH_PING_TEMPLATE = "run util ping -c {count} -W 1 {ip}"
PROBE_ADDRESS = "127.0.0.1"

RAW_EXCERPT_CHARS = 120

_RECEIVED_RE = re.compile(r"(\d+)\s+(?:packets\s+)?received")
_TMSH_ERROR_MARKERS = ("syntax error", "unknown command", "unexpected")


class PingCheckError(F5ClientError):
    """SSH transport or workbook failure during the ping post-process."""


@dataclass
class PingResult:
    status: str  # STATUS_YES / STATUS_NO / STATUS_UNKNOWN / STATUS_NOT_TESTED
    note: str


# ---------------------------------------------------------------------------
# Pure classification and parsing
# ---------------------------------------------------------------------------


def classify_address(raw: Any) -> tuple[str | None, PingResult | None]:
    """(pingable canonical IPv4, None), (None, skip result), or (None, None).

    Only the canonical string returned here is ever interpolated into a
    remote command — this is the command-injection gate.
    """
    text = str(raw).strip() if raw is not None else ""
    if not text:
        return None, None
    if "%" in text:
        # Route domain suffix, checked before parsing so IPv6+rd lands here.
        ip_part, _, route_domain = text.partition("%")
        if not route_domain.isdigit():
            return None, PingResult(STATUS_NOT_TESTED, "FQDN node (not pinged)")
        if route_domain != "0":
            return None, PingResult(STATUS_NOT_TESTED, f"route domain {route_domain} (not pinged)")
        text = ip_part
    try:
        parsed = ipaddress.ip_address(text)
    except ValueError:
        return None, PingResult(STATUS_NOT_TESTED, "FQDN node (not pinged)")
    if parsed.version != 4:
        return None, PingResult(STATUS_NOT_TESTED, "IPv6 address (not pinged)")
    return str(parsed), None


def parse_ping_output(output: str) -> PingResult:
    match = _RECEIVED_RE.search(output or "")
    if not match:
        excerpt = " ".join((output or "").split())[:RAW_EXCERPT_CHARS]
        return PingResult(STATUS_UNKNOWN, f"unparseable output: {excerpt}")
    received = int(match.group(1))
    if received > 0:
        return PingResult(STATUS_YES, f"{received} received")
    return PingResult(STATUS_NO, "0 received")


def looks_like_tmsh_error(output: str) -> bool:
    lowered = (output or "").lower()
    return any(marker in lowered for marker in _TMSH_ERROR_MARKERS)


# ---------------------------------------------------------------------------
# Ping engine (transport-agnostic)
# ---------------------------------------------------------------------------


def _send(run_command: Callable[[str], str], command: str) -> str:
    # Runtime belt-and-braces on top of the fixed templates: this module
    # must be structurally incapable of executing anything but ping.
    if not command.startswith(("ping ", "run util ping ")):
        raise PingCheckError(f"refusing to execute a non-ping command: {command!r}")
    return run_command(command)


def detect_ping_template(run_command: Callable[[str], str], count: int) -> str:
    """One probe against localhost decides bash vs tmsh for the whole run."""
    probe = BASH_PING_TEMPLATE.format(count=1, ip=PROBE_ADDRESS)
    output = _send(run_command, probe)
    if looks_like_tmsh_error(output):
        logger.debug("tmsh login shell detected; using the tmsh ping form")
        return TMSH_PING_TEMPLATE
    return BASH_PING_TEMPLATE


def run_ping_checks(
    addresses: Iterable[str],
    run_command: Callable[[str], str],
    count: int = 2,
    progress: Callable[[str], None] | None = None,
) -> tuple[dict[str, PingResult], str | None]:
    """Ping each unique IPv4 once, sequentially; results keyed by the raw
    address string. On a transport failure mid-run, the results gathered
    so far are kept and the remaining addresses are marked NOT TESTED.
    """
    results: dict[str, PingResult] = {}
    cache: dict[str, PingResult] = {}  # keyed by canonical IPv4
    template: str | None = None
    failure: str | None = None
    pending = list(addresses)
    for index, raw in enumerate(pending):
        ip, skip = classify_address(raw)
        if ip is None:
            if skip is not None:
                results[raw] = skip
            continue
        if ip in cache:
            results[raw] = cache[ip]
            continue
        try:
            if template is None:
                template = detect_ping_template(run_command, count)
            output = _send(run_command, template.format(count=count, ip=ip))
        except PingCheckError as exc:
            failure = str(exc)
            _mark_untested(pending[index:], results, cache)
            break
        result = parse_ping_output(output)
        cache[ip] = result
        results[raw] = result
        if progress:
            progress(f"  ping {ip} ... {result.status} ({result.note})")
    return results, failure


def _mark_untested(
    remaining: Iterable[str], results: dict[str, PingResult], cache: dict[str, PingResult]
) -> None:
    for raw in remaining:
        if raw in results:
            continue
        ip, skip = classify_address(raw)
        if ip is None:
            if skip is not None:
                results[raw] = skip
        elif ip in cache:
            results[raw] = cache[ip]
        else:
            results[raw] = PingResult(STATUS_NOT_TESTED, "not attempted (SSH failure)")


# ---------------------------------------------------------------------------
# SSH transport (the only paramiko code in the package)
# ---------------------------------------------------------------------------


class SSHCommandRunner:
    """One SSH connection to the BIG-IP; one exec channel per command."""

    def __init__(
        self,
        host: str,
        username: str,
        password: str,
        *,
        port: int = 22,
        connect_timeout: int = 15,
        command_timeout: int = 10,
    ):
        self._host = host
        self._username = username
        self._password = password
        self._port = port
        self._connect_timeout = connect_timeout
        self._command_timeout = command_timeout
        self._client: Any = None
        self._paramiko: Any = None

    def connect(self) -> None:
        try:
            import paramiko
        except ImportError as exc:
            raise PingCheckError(
                "The 'ping' command needs paramiko. Install it with: pip install 'f5audit[ssh]'"
            ) from exc
        self._paramiko = paramiko
        client = paramiko.SSHClient()
        # Management-network tool: auto-accept unknown host keys, in the
        # same spirit as --insecure for TLS (documented in the README).
        client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
        try:
            client.connect(
                self._host,
                port=self._port,
                username=self._username,
                password=self._password,
                timeout=self._connect_timeout,
                look_for_keys=True,
                allow_agent=True,
            )
        except paramiko.AuthenticationException as exc:
            raise PingCheckError(
                f"SSH authentication failed for {self._username}@{self._host}"
            ) from exc
        except (paramiko.SSHException, OSError) as exc:
            raise PingCheckError(
                f"Could not open an SSH connection to {self._host}:{self._port}: {exc}"
            ) from exc
        self._client = client

    def run(self, command: str) -> str:
        if self._client is None:
            raise PingCheckError("SSH connection is not open")
        try:
            _, stdout, stderr = self._client.exec_command(command, timeout=self._command_timeout)
            # tmsh errors may arrive on either stream; the parser sees both.
            return stdout.read().decode(errors="replace") + stderr.read().decode(errors="replace")
        except (self._paramiko.SSHException, OSError) as exc:
            raise PingCheckError(f"SSH command failed on {self._host}: {exc}") from exc

    def close(self) -> None:
        if self._client is not None:
            self._client.close()
            self._client = None

    def __enter__(self) -> SSHCommandRunner:
        self.connect()
        return self

    def __exit__(self, exc_type, exc_value, traceback) -> None:
        self.close()


# ---------------------------------------------------------------------------
# Workbook I/O
# ---------------------------------------------------------------------------


def load_report(path: str):
    try:
        return load_workbook(path)
    except Exception as exc:  # noqa: BLE001 - surfaced as a friendly error
        raise PingCheckError(f"Could not open report {path}: {exc}") from exc


def _locate_column(sheet, header: str) -> int | None:
    for cell in sheet[1]:
        if str(cell.value) == header:
            return cell.column
    return None


def _cell_text(sheet, row: int, column: int) -> str:
    value = sheet.cell(row=row, column=column).value
    return str(value).strip() if value is not None else ""


def collect_addresses(workbook) -> tuple[list[str], list[str]]:
    """Ordered unique raw addresses from the Orphan Nodes IP column."""
    warnings: list[str] = []
    if ADDRESS_SOURCE_SHEET not in workbook.sheetnames:
        warnings.append(f"Sheet '{ADDRESS_SOURCE_SHEET}' not found in the report")
        return [], warnings
    sheet = workbook[ADDRESS_SOURCE_SHEET]
    ip_column = _locate_column(sheet, IP_HEADER)
    if ip_column is None:
        warnings.append(f"No '{IP_HEADER}' column on sheet '{ADDRESS_SOURCE_SHEET}'")
        return [], warnings
    addresses: list[str] = []
    seen: set[str] = set()
    for row in range(2, sheet.max_row + 1):
        text = _cell_text(sheet, row, ip_column)
        if text and text not in seen:
            seen.add(text)
            addresses.append(text)
    return addresses, warnings


def _write_header(sheet, column: int, header: str) -> None:
    cell = sheet.cell(row=1, column=column, value=header)
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
    cell.alignment = Alignment(vertical="center")


def enrich_workbook(workbook, results: dict[str, PingResult]) -> list[str]:
    """Write the ping columns into every enriched sheet, in place.

    Idempotent: an existing 'Ping (from F5)' column is overwritten, so a
    re-run never appends duplicate columns. Rows whose IP was not pinged
    (IN USE nodes, empty cells) get blank cells.
    """
    warnings: list[str] = []
    for title in ENRICHED_SHEETS:
        if title not in workbook.sheetnames:
            warnings.append(f"Sheet '{title}' not found; skipped")
            continue
        sheet = workbook[title]
        ip_column = _locate_column(sheet, IP_HEADER)
        if ip_column is None:
            warnings.append(f"No '{IP_HEADER}' column on sheet '{title}'; skipped")
            continue
        status_column = _locate_column(sheet, STATUS_HEADER)
        if status_column is None:
            status_column = sheet.max_column + 1
        note_column = status_column + 1
        _write_header(sheet, status_column, STATUS_HEADER)
        _write_header(sheet, note_column, NOTE_HEADER)
        status_width = len(STATUS_HEADER)
        note_width = len(NOTE_HEADER)
        for row in range(2, sheet.max_row + 1):
            result = results.get(_cell_text(sheet, row, ip_column))
            status = result.status if result else ""
            note = result.note if result else ""
            sheet.cell(row=row, column=status_column, value=status)
            sheet.cell(row=row, column=note_column, value=note)
            status_width = max(status_width, len(status))
            note_width = max(note_width, len(note))
        last_column = get_column_letter(sheet.max_column)
        sheet.auto_filter.ref = f"A1:{last_column}{max(sheet.max_row, 1)}"
        sheet.column_dimensions[get_column_letter(status_column)].width = min(
            status_width + 2, MAX_COLUMN_WIDTH
        )
        sheet.column_dimensions[get_column_letter(note_column)].width = min(
            note_width + 2, MAX_COLUMN_WIDTH
        )
    return warnings
