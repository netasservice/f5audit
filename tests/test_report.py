"""Report tests: table content, Excel structure, CSV output."""

import csv

from openpyxl import load_workbook

from f5audit.analyzer import Analyzer, Verdict
from f5audit.correlator import correlate
from f5audit.models import IRule, PoolMember
from f5audit.parsing import parse_collection
from f5audit.report import build_tables, default_report_name, write_csv, write_xlsx
from tests.conftest import build_collection


def make_tables():
    parsed = parse_collection(build_collection())
    correlation = correlate(parsed)
    analysis = Analyzer(parsed, correlation).run()
    return parsed, build_tables(parsed, correlation, analysis)


def test_expected_sheets_exist():
    _, tables = make_tables()
    assert list(tables) == [
        "summary",
        "inventory",
        "orphan_nodes",
        "pools",
        "inactive_virtuals",
        "dead_chains",
        "orphan_monitors",
        "manual_review",
    ]


def test_inventory_has_member_rows_and_orphan_node_rows():
    _, tables = make_tables()
    inventory = tables["inventory"]
    first_column = [row[0] for row in inventory.rows]
    assert "/Common/node-web-1" in first_column  # pool member row
    assert "/Common/node-orphan" in first_column  # node without pool

    member_row = next(
        r for r in inventory.rows if r[0] == "/Common/node-web-1" and r[5] == "/Common/pool-web"
    )
    assert member_row[1] == "10.0.0.1"
    assert member_row[4] == Verdict.IN_USE  # node verdict
    assert member_row[6] == "80"
    assert "/Common/vs-web" in member_row[10]
    assert member_row[14] == "/Common/irule-static"  # iRules attached to the VS
    assert member_row[17] == Verdict.IN_USE  # pool verdict

    orphan_row = next(r for r in inventory.rows if r[0] == "/Common/node-orphan")
    assert orphan_row[4] == Verdict.ORPHAN
    assert orphan_row[17] == Verdict.ORPHAN


def test_orphan_sheets_only_contain_non_in_use_objects():
    _, tables = make_tables()
    assert [row[0] for row in tables["orphan_nodes"].rows] == [
        "/Common/node-dead",
        "/Common/node-orphan",
    ]
    pool_names = [row[0] for row in tables["pools"].rows]
    assert "/Common/pool-orphan" in pool_names
    assert "/Common/pool-idle" in pool_names
    assert "/Common/pool-dead" in pool_names
    assert "/Common/pool-web" not in pool_names
    monitor_names = [row[0] for row in tables["orphan_monitors"].rows]
    assert monitor_names == ["/Common/mon-orphan"]


def test_suggested_commands_only_for_orphans():
    _, tables = make_tables()
    for row in tables["pools"].rows:
        command = row[-1]
        if row[8] == Verdict.ORPHAN:
            assert command.startswith("tmsh delete ltm pool ")
        else:
            # OFFLINE candidates included: their commands live only on the
            # Dead Chains sheet.
            assert command == ""


def test_dead_chains_sheet_groups_the_whole_chain():
    _, tables = make_tables()
    rows = tables["dead_chains"].rows
    assert [row[0] for row in rows] == ["/Common/pool-dead"]
    row = rows[0]
    assert row[3] == "/Common/node-dead:443"
    assert row[6] == "/Common/vs-dead"
    assert row[9] == Verdict.OFFLINE_CANDIDATE
    commands = row[-1].splitlines()
    assert commands == [
        "tmsh delete ltm virtual /Common/vs-dead",
        "tmsh delete ltm pool /Common/pool-dead",
        "tmsh delete ltm node /Common/node-dead",
    ]


def test_dead_chains_sheet_omits_node_command_when_alive_elsewhere():
    parsed = parse_collection(build_collection())
    # node-dead is also an available member of pool-web.
    parsed.pools["/Common/pool-web"].members.append(
        PoolMember(
            node_full_path="/Common/node-dead",
            port="80",
            partition="Common",
            admin_state="monitor-enabled",
            availability="available",
        )
    )
    correlation = correlate(parsed)
    analysis = Analyzer(parsed, correlation).run()
    tables = build_tables(parsed, correlation, analysis)
    row = tables["dead_chains"].rows[0]
    commands = row[-1].splitlines()
    assert "tmsh delete ltm node /Common/node-dead" not in commands
    assert "tmsh delete ltm pool /Common/pool-dead" in commands


def test_dead_chains_sheet_keeps_capped_pool_without_pool_command():
    parsed = parse_collection(build_collection())
    parsed.irules["/Common/irule-dyn"] = IRule(
        full_path="/Common/irule-dyn",
        partition="Common",
        name="irule-dyn",
        definition="pool $x",
        has_dynamic_pool_selection=True,
    )
    parsed.virtuals["/Common/vs-web"].irules.append("/Common/irule-dyn")
    correlation = correlate(parsed)
    analysis = Analyzer(parsed, correlation).run()
    tables = build_tables(parsed, correlation, analysis)
    rows = tables["dead_chains"].rows
    assert [row[0] for row in rows] == ["/Common/pool-dead"]
    row = rows[0]
    assert row[9] == Verdict.MANUAL_REVIEW
    assert row[5] == Verdict.OFFLINE_CANDIDATE  # node verdict
    commands = row[-1].splitlines()
    assert commands == [
        "tmsh delete ltm virtual /Common/vs-dead",
        "tmsh delete ltm node /Common/node-dead",
    ]


def test_inventory_shows_node_and_pool_verdicts_separately():
    """A dead node in a pool capped by dynamic iRules: node OFFLINE, pool
    MANUAL REVIEW, both visible on the same row."""
    parsed = parse_collection(build_collection())
    parsed.irules["/Common/irule-dyn"] = IRule(
        full_path="/Common/irule-dyn",
        partition="Common",
        name="irule-dyn",
        definition="pool $x",
        has_dynamic_pool_selection=True,
    )
    parsed.virtuals["/Common/vs-web"].irules.append("/Common/irule-dyn")
    correlation = correlate(parsed)
    analysis = Analyzer(parsed, correlation).run()
    tables = build_tables(parsed, correlation, analysis)
    row = next(r for r in tables["inventory"].rows if r[0] == "/Common/node-dead")
    assert row[4] == Verdict.OFFLINE_CANDIDATE
    assert row[17] == Verdict.MANUAL_REVIEW


def test_network_columns_are_appended_to_inventory_and_orphan_nodes():
    _, tables = make_tables()
    inventory = tables["inventory"]
    assert inventory.headers[-2:] == ["ARP MAC", "Network note"]
    assert inventory.verdict_columns == (4, 17)  # unchanged by the appended columns

    member_row = next(
        r for r in inventory.rows if r[0] == "/Common/node-web-1" and r[5] == "/Common/pool-web"
    )
    assert member_row[-2] == "00:00:5e:00:53:01"
    assert member_row[-1] == "in ARP table (MAC 00:00:5e:00:53:01)"

    orphan_nodes = tables["orphan_nodes"]
    assert orphan_nodes.headers[-2:] == ["ARP MAC", "Network note"]
    assert orphan_nodes.verdict_columns == (4,)
    rows = {row[0]: row for row in orphan_nodes.rows}
    # /26 self-IP: node-dead (10.0.0.50) is local, node-orphan (10.0.0.99) routed.
    assert rows["/Common/node-dead"][-1] == "on local subnet, no ARP entry (idle or down)"
    assert rows["/Common/node-orphan"][-1] == "not directly connected (behind a router)"


def test_network_columns_degrade_on_old_raw_cache():
    parsed = parse_collection(build_collection(network=False))
    correlation = correlate(parsed)
    analysis = Analyzer(parsed, correlation).run()
    tables = build_tables(parsed, correlation, analysis)
    row = next(r for r in tables["inventory"].rows if r[0] == "/Common/node-web-1")
    assert row[-2:] == ["", "network data not collected"]
    summary = {str(r[0]): r[1] for r in tables["summary"].rows}
    assert summary["Network data (ARP/self-IP)"] == "not collected (old cache)"


def test_network_note_carries_standby_annotation():
    parsed = parse_collection(build_collection(standby=True))
    correlation = correlate(parsed)
    analysis = Analyzer(parsed, correlation).run()
    tables = build_tables(parsed, correlation, analysis)
    row = next(r for r in tables["inventory"].rows if r[0] == "/Common/node-web-1")
    assert row[-1].endswith("[standby unit: ARP reflects this unit only]")


def test_summary_reports_resumed_collection():
    data = build_collection()
    data.meta["resumed_at"] = ["2026-08-26T10:00:00+00:00"]
    parsed = parse_collection(data)
    correlation = correlate(parsed)
    analysis = Analyzer(parsed, correlation).run()
    tables = build_tables(parsed, correlation, analysis)
    summary = {str(r[0]): r[1] for r in tables["summary"].rows}
    assert summary["Collection resumed (mixed timestamps)"] == "2026-08-26T10:00:00+00:00"


def test_summary_contains_system_info_and_counts():
    _, tables = make_tables()
    rows = {str(row[0]): row[1] for row in tables["summary"].rows}
    assert rows["Hostname"] == "bigip1.example.net"
    assert rows["HA state"] == "active"
    assert Verdict.ORPHAN in rows


def test_write_xlsx(tmp_path):
    _, tables = make_tables()
    out = tmp_path / "report.xlsx"
    write_xlsx(tables, str(out))

    workbook = load_workbook(str(out))
    assert "Summary" in workbook.sheetnames
    assert "Inventory" in workbook.sheetnames

    inventory = workbook["Inventory"]
    assert inventory.freeze_panes == "A2"
    assert inventory.auto_filter.ref is not None
    assert inventory.cell(row=1, column=1).value == "Node"
    assert inventory.max_row > 1

    # Verdict cells carry the conditional fill colors, in both the node
    # verdict and the pool verdict columns.
    node_verdict_column, pool_verdict_column = (c + 1 for c in tables["inventory"].verdict_columns)
    fills = set()
    node_fills = set()
    for row_index in range(2, inventory.max_row + 1):
        cell = inventory.cell(row=row_index, column=pool_verdict_column)
        if cell.fill and cell.fill.fgColor and cell.fill.fgColor.rgb:
            fills.add(cell.fill.fgColor.rgb)
        node_cell = inventory.cell(row=row_index, column=node_verdict_column)
        if node_cell.fill and node_cell.fill.fgColor and node_cell.fill.fgColor.rgb:
            node_fills.add(node_cell.fill.fgColor.rgb)
    assert "00C6EFCE" in fills or "FFC6EFCE" in fills  # green for IN USE
    assert "00FFC7CE" in node_fills or "FFFFC7CE" in node_fills  # red for ORPHAN node

    dead_chains = workbook["Dead Chains"]
    (verdict_column,) = (c + 1 for c in tables["dead_chains"].verdict_columns)
    cell = dead_chains.cell(row=2, column=verdict_column)
    assert cell.fill.fgColor.rgb in ("00CCC0DA", "FFCCC0DA")  # purple for OFFLINE


def test_write_csv(tmp_path):
    _, tables = make_tables()
    written = write_csv(tables, str(tmp_path))
    assert len(written) == 8

    with open(tmp_path / "inventory.csv", encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    assert rows[0][0] == "Node"
    assert len(rows) > 1


def test_default_report_name():
    name = default_report_name("bigip1.example.net")
    assert name.startswith("f5audit_bigip1.example.net_")
    assert name.endswith(".xlsx")
    assert default_report_name("host", "csv").endswith(
        ("0", "1", "2", "3", "4", "5", "6", "7", "8", "9")
    )
