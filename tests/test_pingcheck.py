"""Ping post-process tests. No test opens a socket or needs paramiko."""

import re
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from f5audit.analyzer import Analyzer
from f5audit.cli import EXIT_ERROR, main
from f5audit.correlator import correlate
from f5audit.parsing import parse_collection
from f5audit.pingcheck import (
    BASH_PING_TEMPLATE,
    STATUS_NO,
    STATUS_NOT_TESTED,
    STATUS_UNKNOWN,
    STATUS_YES,
    TMSH_PING_TEMPLATE,
    PingCheckError,
    PingResult,
    SSHCommandRunner,
    classify_address,
    collect_addresses,
    enrich_workbook,
    load_report,
    parse_ping_output,
    run_ping_checks,
)
from f5audit.report import build_tables, write_xlsx
from tests.conftest import build_collection

PING_COMMAND_RE = re.compile(r"^(ping|run util ping) -c \d+ -W 1 (\d{1,3}\.){3}\d{1,3}$")


def make_report(tmp_path) -> Path:
    parsed = parse_collection(build_collection())
    correlation = correlate(parsed)
    analysis = Analyzer(parsed, correlation).run()
    path = tmp_path / "report.xlsx"
    write_xlsx(build_tables(parsed, correlation, analysis), str(path))
    return path


class FakeRunner:
    """Stands in for SSHCommandRunner.run; serves canned ping output."""

    def __init__(self, shell="bash", reachable=frozenset(), fail_from=None):
        self.shell = shell
        self.reachable = set(reachable)
        self.fail_from = fail_from  # 1-based command index to start failing at
        self.commands = []

    def __call__(self, command):
        self.commands.append(command)
        if self.fail_from is not None and len(self.commands) >= self.fail_from:
            raise PingCheckError("connection lost")
        if self.shell == "tmsh" and command.startswith("ping "):
            return 'Syntax Error: unexpected argument "ping"'
        ip = command.split()[-1]
        received = 2 if ip in self.reachable else 0
        return f"2 packets transmitted, {received} received, 0% packet loss"


# ---------------------------------------------------------------------------
# Classification
# ---------------------------------------------------------------------------


def test_classify_address_pingable_forms():
    assert classify_address("10.0.0.1") == ("10.0.0.1", None)
    assert classify_address(" 10.0.0.1 ") == ("10.0.0.1", None)
    assert classify_address("10.0.0.1%0") == ("10.0.0.1", None)


def test_classify_address_skips():
    ip, skip = classify_address("10.0.0.1%2")
    assert ip is None and skip.note == "route domain 2 (not pinged)"
    ip, skip = classify_address("2001:db8::1")
    assert ip is None and skip.note == "IPv6 address (not pinged)"
    # Route domain is checked before parsing, so IPv6+rd lands there.
    ip, skip = classify_address("fe80::1%3")
    assert ip is None and skip.note == "route domain 3 (not pinged)"
    ip, skip = classify_address("app.example.net")
    assert ip is None and skip.note == "FQDN node (not pinged)"
    ip, skip = classify_address("10.0.0.1%abc")
    assert ip is None and skip.status == STATUS_NOT_TESTED


def test_classify_address_empty_forms():
    assert classify_address("") == (None, None)
    assert classify_address(None) == (None, None)
    assert classify_address("   ") == (None, None)


def test_classify_address_rejects_injection_attempts():
    for hostile in ("10.0.0.1; reboot", "$(id)", "10.0.0.1 && ls", "10.0.0.1 -c 100000"):
        ip, skip = classify_address(hostile)
        assert ip is None
        assert skip.status == STATUS_NOT_TESTED


# ---------------------------------------------------------------------------
# Output parsing
# ---------------------------------------------------------------------------


def test_parse_ping_output():
    assert parse_ping_output("5 packets transmitted, 2 received, 60% loss").status == STATUS_YES
    assert parse_ping_output("2 packets transmitted, 2 packets received").status == STATUS_YES
    assert parse_ping_output("2 packets transmitted, 0 received, 100% loss").status == STATUS_NO
    garbled = parse_ping_output("PING statistics garbled " + "x" * 300)
    assert garbled.status == STATUS_UNKNOWN
    assert garbled.note.startswith("unparseable output: ")
    assert len(garbled.note) <= len("unparseable output: ") + 120


# ---------------------------------------------------------------------------
# Engine: guard, shell detection, dedupe, degradation
# ---------------------------------------------------------------------------


def test_command_templates_are_ping_only():
    assert BASH_PING_TEMPLATE.startswith("ping ")
    assert TMSH_PING_TEMPLATE.startswith("run util ping ")


def test_engine_refuses_non_ping_commands():
    from f5audit.pingcheck import _send

    with pytest.raises(PingCheckError, match="non-ping"):
        _send(lambda cmd: "", "tmsh delete ltm node /Common/x")


def test_run_ping_checks_only_emits_ping_commands():
    runner = FakeRunner(reachable={"10.0.0.1"})
    run_ping_checks(["10.0.0.1", "10.0.0.2", "bad; rm -rf /"], runner)
    assert runner.commands  # probe + 2 pings
    for command in runner.commands:
        assert PING_COMMAND_RE.match(command), command


def test_shell_detection_bash():
    runner = FakeRunner(shell="bash", reachable={"10.0.0.1"})
    results, failure = run_ping_checks(["10.0.0.1"], runner)
    assert failure is None
    assert runner.commands[0] == "ping -c 1 -W 1 127.0.0.1"  # single probe
    assert runner.commands[1] == "ping -c 2 -W 1 10.0.0.1"
    assert results["10.0.0.1"].status == STATUS_YES


def test_shell_detection_tmsh():
    runner = FakeRunner(shell="tmsh", reachable={"10.0.0.1"})
    results, failure = run_ping_checks(["10.0.0.1", "10.0.0.2"], runner)
    assert failure is None
    # One probe, then the tmsh form for every real ping.
    assert runner.commands[0].startswith("ping ")
    assert runner.commands[1] == "run util ping -c 2 -W 1 10.0.0.1"
    assert runner.commands[2] == "run util ping -c 2 -W 1 10.0.0.2"
    assert results["10.0.0.1"].status == STATUS_YES
    assert results["10.0.0.2"].status == STATUS_NO


def test_run_ping_checks_deduplicates_addresses():
    runner = FakeRunner(reachable={"10.0.0.1"})
    results, _ = run_ping_checks(["10.0.0.1", "10.0.0.1%0", "10.0.0.1"], runner)
    pings = [c for c in runner.commands if not c.endswith("127.0.0.1")]
    assert pings == ["ping -c 2 -W 1 10.0.0.1"]
    assert results["10.0.0.1"].status == STATUS_YES
    assert results["10.0.0.1%0"].status == STATUS_YES


def test_run_ping_checks_keeps_partial_results_on_ssh_failure():
    # Probe + first ping succeed; the second ping dies.
    runner = FakeRunner(reachable={"10.0.0.1"}, fail_from=3)
    addresses = ["10.0.0.1", "10.0.0.2", "10.0.0.3", "host.example.net"]
    results, failure = run_ping_checks(addresses, runner)
    assert failure == "connection lost"
    assert results["10.0.0.1"].status == STATUS_YES  # kept
    assert results["10.0.0.2"] == PingResult(STATUS_NOT_TESTED, "not attempted (SSH failure)")
    assert results["10.0.0.3"] == PingResult(STATUS_NOT_TESTED, "not attempted (SSH failure)")
    assert results["host.example.net"].note == "FQDN node (not pinged)"  # skip still classified


# ---------------------------------------------------------------------------
# Workbook: address source, enrichment, idempotency, degradation
# ---------------------------------------------------------------------------


def test_collect_addresses_excludes_in_use_nodes(tmp_path):
    workbook = load_report(str(make_report(tmp_path)))
    addresses, warnings = collect_addresses(workbook)
    assert warnings == []
    # Orphan Nodes only: node-dead + node-orphan. The IN USE node-web-1
    # (10.0.0.1) must never be pinged.
    assert addresses == ["10.0.0.50", "10.0.0.99"]


def test_enrich_workbook_replicates_results_across_sheets(tmp_path):
    path = make_report(tmp_path)
    workbook = load_report(str(path))
    addresses, _ = collect_addresses(workbook)
    runner = FakeRunner(reachable={"10.0.0.99"})
    results, failure = run_ping_checks(addresses, runner)
    assert failure is None
    assert enrich_workbook(workbook, results) == []
    workbook.save(str(path))

    reloaded = load_workbook(str(path))
    for title in ("Inventory", "Orphan Nodes"):
        sheet = reloaded[title]
        headers = [cell.value for cell in sheet[1]]
        assert headers[-2:] == ["Ping (from F5)", "Ping note"]
        header_cell = sheet.cell(row=1, column=len(headers) - 1)
        assert header_cell.font.bold
        assert header_cell.fill.fgColor.rgb.endswith("D9D9D9")
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref.endswith(f"{sheet.max_row}")
        by_ip = {}
        for row in range(2, sheet.max_row + 1):
            ip = sheet.cell(row=row, column=2).value
            status = sheet.cell(row=row, column=len(headers) - 1).value
            by_ip.setdefault(ip, set()).add(status)
        # Replicated on every row sharing the IP, on both sheets.
        assert by_ip.get("10.0.0.50") == {STATUS_NO}
        assert by_ip.get("10.0.0.99") == {STATUS_YES}
    inventory = reloaded["Inventory"]
    for row in range(2, inventory.max_row + 1):
        ip = inventory.cell(row=row, column=2).value
        status = inventory.cell(row=row, column=inventory.max_column - 1).value
        if ip == "10.0.0.1":  # IN USE: never pinged, cells stay blank
            assert status in (None, "")
        if ip in (None, ""):  # "(no members)" rows
            assert status in (None, "")


def test_enrich_workbook_is_idempotent(tmp_path):
    path = make_report(tmp_path)
    workbook = load_report(str(path))
    addresses, _ = collect_addresses(workbook)
    first, _ = run_ping_checks(addresses, FakeRunner(reachable={"10.0.0.50", "10.0.0.99"}))
    enrich_workbook(workbook, first)
    workbook.save(str(path))

    workbook = load_report(str(path))
    before_columns = workbook["Inventory"].max_column
    second, _ = run_ping_checks(addresses, FakeRunner(reachable=set()))
    enrich_workbook(workbook, second)
    workbook.save(str(path))

    reloaded = load_workbook(str(path))
    sheet = reloaded["Orphan Nodes"]
    headers = [cell.value for cell in sheet[1]]
    assert headers.count("Ping (from F5)") == 1
    assert reloaded["Inventory"].max_column == before_columns
    statuses = {
        sheet.cell(row=row, column=len(headers) - 1).value for row in range(2, sheet.max_row + 1)
    }
    assert statuses == {STATUS_NO}  # overwritten, not duplicated


def test_enrich_workbook_degrades_when_a_sheet_is_missing(tmp_path):
    path = make_report(tmp_path)
    workbook = load_report(str(path))
    addresses, _ = collect_addresses(workbook)
    results, _ = run_ping_checks(addresses, FakeRunner())
    del workbook["Inventory"]
    warnings = enrich_workbook(workbook, results)
    assert any("Inventory" in warning for warning in warnings)
    headers = [cell.value for cell in workbook["Orphan Nodes"][1]]
    assert headers[-2:] == ["Ping (from F5)", "Ping note"]


def test_collect_addresses_warns_without_orphan_nodes_sheet(tmp_path):
    workbook = load_report(str(make_report(tmp_path)))
    del workbook["Orphan Nodes"]
    addresses, warnings = collect_addresses(workbook)
    assert addresses == []
    assert any("Orphan Nodes" in warning for warning in warnings)


# ---------------------------------------------------------------------------
# Lazy paramiko
# ---------------------------------------------------------------------------


def test_importing_the_module_does_not_import_paramiko():
    code = "import sys, f5audit.pingcheck; sys.exit(1 if 'paramiko' in sys.modules else 0)"
    completed = subprocess.run(
        [sys.executable, "-c", code],
        cwd=str(Path(__file__).resolve().parent.parent),
    )
    assert completed.returncode == 0


def test_missing_paramiko_gives_install_hint(monkeypatch):
    monkeypatch.setitem(sys.modules, "paramiko", None)
    runner = SSHCommandRunner("192.0.2.1", "auditor", "secret")
    with pytest.raises(PingCheckError, match=r"pip install 'f5audit\[ssh\]'"):
        runner.connect()


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def test_cmd_ping_rejects_csv_report_without_prompting(tmp_path, capsys):
    assert main(["ping", "--report", str(tmp_path / "report.csv")]) == EXIT_ERROR
    assert "xlsx reports only" in capsys.readouterr().err


def test_cmd_ping_rejects_missing_report(tmp_path, capsys):
    assert main(["ping", "--report", str(tmp_path / "missing.xlsx")]) == EXIT_ERROR
    assert "not found" in capsys.readouterr().err


def test_cmd_ping_requires_host_after_validating_report(tmp_path, capsys):
    path = make_report(tmp_path)
    assert main(["ping", "--report", str(path)]) == EXIT_ERROR
    assert "--host is required" in capsys.readouterr().err
