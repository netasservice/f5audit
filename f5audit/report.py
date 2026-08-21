"""Report generation: multi-sheet Excel workbook (openpyxl) or CSV set.

The "suggested command" columns are informational text for the human
running the change control; this tool never executes anything.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .analyzer import AnalysisResult, Verdict
from .correlator import Correlation
from .parsing import ParsedData

VERDICT_FILLS = {
    Verdict.ORPHAN: "FFC7CE",  # red
    Verdict.MANUAL_REVIEW: "FFEB9C",  # yellow
    Verdict.UNRELIABLE_STANDBY: "FFEB9C",  # yellow
    Verdict.UNRELIABLE_INVENTORY: "FFEB9C",
    Verdict.INACTIVE: "FCD5B4",  # orange
    Verdict.OFFLINE_CANDIDATE: "CCC0DA",  # purple
    Verdict.IN_USE: "C6EFCE",  # green
}

HEADER_FILL = "D9D9D9"
MAX_COLUMN_WIDTH = 60


@dataclass
class ReportTable:
    title: str
    headers: list[str]
    rows: list[list[object]] = field(default_factory=list)
    verdict_column: int | None = None  # 0-based index into headers


def default_report_name(hostname: str, fmt: str = "xlsx") -> str:
    safe_host = re.sub(r"[^A-Za-z0-9._-]", "_", hostname or "unknown")
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    suffix = "" if fmt == "csv" else ".xlsx"
    return f"f5audit_{safe_host}_{stamp}{suffix}"


# ---------------------------------------------------------------------------
# Table construction
# ---------------------------------------------------------------------------


def _join(values) -> str:
    return ", ".join(sorted(values)) if values else ""


def build_tables(
    parsed: ParsedData, correlation: Correlation, analysis: AnalysisResult
) -> dict[str, ReportTable]:
    return {
        "summary": _build_summary(parsed, analysis),
        "inventory": _build_inventory(parsed, correlation, analysis),
        "orphan_nodes": _build_orphan_nodes(parsed, analysis),
        "pools": _build_pools(parsed, correlation, analysis),
        "inactive_virtuals": _build_inactive_virtuals(parsed, analysis),
        "dead_chains": _build_dead_chains(parsed, correlation, analysis),
        "orphan_monitors": _build_orphan_monitors(parsed, analysis),
        "manual_review": _build_manual_review(analysis),
    }


def _build_summary(parsed: ParsedData, analysis: AnalysisResult) -> ReportTable:
    system = parsed.system
    table = ReportTable("Summary", ["Item", "Value"])
    table.rows = [
        ["Hostname", system.hostname],
        ["BIG-IP version", system.version],
        ["HA state", system.failover_state],
        ["Active device", system.active_device],
        ["Uptime context", system.uptime],
        ["Collection timestamp (UTC)", system.collection_timestamp],
        ["Partitions collected", _join(system.partitions_collected)],
        ["Partitions denied", _join(system.partitions_denied)],
        ["Missing endpoints", _join(system.missing_endpoints)],
        ["Traffic analysis skipped", "YES" if analysis.stats_analysis_skipped else "no"],
        ["", ""],
        ["Objects", ""],
        ["Nodes", len(parsed.nodes)],
        ["Pools", len(parsed.pools)],
        ["Virtual servers", len(parsed.virtuals)],
        ["iRules", len(parsed.irules)],
        ["Policies", len(parsed.policies)],
        ["Monitors", len(parsed.monitors)],
        ["", ""],
        ["Verdict counts", ""],
    ]
    for verdict, count in sorted(analysis.verdict_counts().items()):
        table.rows.append([verdict, count])
    if analysis.warnings:
        table.rows.append(["", ""])
        table.rows.append(["WARNINGS", ""])
        for warning in analysis.warnings:
            table.rows.append(["!", warning])
    return table


def _vs_summary(parsed: ParsedData, vs_paths) -> tuple:
    """(names, destinations, states, total_conns, attached iRules) joined
    for a VS set."""
    names, destinations, states, conns, irules = [], [], [], [], []
    for path in sorted(vs_paths):
        virtual = parsed.virtuals.get(path)
        if not virtual:
            names.append(path)
            continue
        names.append(path)
        destinations.append(virtual.destination)
        states.append(f"{virtual.admin_state}/{virtual.availability or '?'}")
        if virtual.total_conns is not None:
            conns.append(str(virtual.total_conns))
        irules.extend(irule for irule in virtual.irules if irule not in irules)
    return (
        ", ".join(names),
        ", ".join(destinations),
        ", ".join(states),
        ", ".join(conns),
        ", ".join(irules),
    )


def _build_inventory(
    parsed: ParsedData, correlation: Correlation, analysis: AnalysisResult
) -> ReportTable:
    headers = [
        "Node",
        "IP",
        "Partition",
        "Node status",
        "Pool",
        "Port",
        "Member status",
        "Effective monitor",
        "LB method",
        "Virtual servers",
        "VIP:Port",
        "VS status",
        "VS total conns",
        "VS iRules",
        "iRules selecting pool",
        "Policies forwarding to pool",
        "Verdict",
        "Notes",
    ]
    table = ReportTable("Inventory", headers, verdict_column=16)

    nodes_in_pools = set()
    for pool_path, pool in sorted(parsed.pools.items()):
        verdict = analysis.pool_verdicts.get(pool_path)
        vs_paths = correlation.pool_to_virtuals.get(pool_path, set())
        vs_names, vips, vs_states, vs_conns, vs_irules = _vs_summary(parsed, vs_paths)
        irule_refs = _join(correlation.pool_to_irules.get(pool_path, set()))
        policy_refs = _join(correlation.pool_to_policies.get(pool_path, set()))
        members = pool.members or [None]
        for member in members:
            node = parsed.nodes.get(member.node_full_path) if member else None
            if member:
                nodes_in_pools.add(member.node_full_path)
            monitor = _join(pool.monitors) or (node.monitor if node else "")
            table.rows.append(
                [
                    member.node_full_path if member else "(no members)",
                    node.address if node else "",
                    pool.partition,
                    f"{node.admin_state}/{node.availability or '?'}" if node else "",
                    pool_path,
                    member.port if member else "",
                    f"{member.admin_state}/{member.availability or '?'}" if member else "",
                    monitor,
                    pool.lb_method,
                    vs_names,
                    vips,
                    vs_states,
                    vs_conns,
                    vs_irules,
                    irule_refs,
                    policy_refs,
                    verdict.verdict if verdict else "",
                    verdict.notes if verdict else "",
                ]
            )

    # Nodes that belong to no pool get their own rows.
    for node_path, node in sorted(parsed.nodes.items()):
        if node_path in nodes_in_pools:
            continue
        verdict = analysis.node_verdicts.get(node_path)
        table.rows.append(
            [
                node_path,
                node.address,
                node.partition,
                f"{node.admin_state}/{node.availability or '?'}",
                "",
                "",
                "",
                node.monitor,
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                verdict.verdict if verdict else "",
                verdict.notes if verdict else "",
            ]
        )
    return table


def _build_orphan_nodes(parsed: ParsedData, analysis: AnalysisResult) -> ReportTable:
    headers = [
        "Node",
        "IP",
        "Partition",
        "Monitor",
        "Verdict",
        "Notes",
        "Suggested command (informational)",
    ]
    table = ReportTable("Orphan Nodes", headers, verdict_column=4)
    for path, verdict in sorted(analysis.node_verdicts.items()):
        if verdict.verdict == Verdict.IN_USE:
            continue
        node = parsed.nodes[path]
        table.rows.append(
            [
                path,
                node.address,
                node.partition,
                node.monitor,
                verdict.verdict,
                verdict.notes,
                f"tmsh delete ltm node {path}" if verdict.verdict == Verdict.ORPHAN else "",
            ]
        )
    return table


def _build_pools(
    parsed: ParsedData, correlation: Correlation, analysis: AnalysisResult
) -> ReportTable:
    headers = [
        "Pool",
        "Partition",
        "LB method",
        "Members",
        "Monitors",
        "Virtual servers",
        "iRules selecting pool",
        "Policies forwarding to pool",
        "Verdict",
        "Notes",
        "Suggested command (informational)",
    ]
    table = ReportTable("Orphan-Inactive Pools", headers, verdict_column=8)
    for path, verdict in sorted(analysis.pool_verdicts.items()):
        if verdict.verdict == Verdict.IN_USE:
            continue
        pool = parsed.pools[path]
        table.rows.append(
            [
                path,
                pool.partition,
                pool.lb_method,
                ", ".join(f"{m.node_full_path}:{m.port}" for m in pool.members),
                _join(pool.monitors),
                _join(correlation.pool_to_virtuals.get(path, set())),
                _join(correlation.pool_to_irules.get(path, set())),
                _join(correlation.pool_to_policies.get(path, set())),
                verdict.verdict,
                verdict.notes,
                f"tmsh delete ltm pool {path}" if verdict.verdict == Verdict.ORPHAN else "",
            ]
        )
    return table


def _build_inactive_virtuals(parsed: ParsedData, analysis: AnalysisResult) -> ReportTable:
    headers = [
        "Virtual server",
        "VIP:Port",
        "Default pool",
        "iRules",
        "State",
        "Availability",
        "Total conns",
        "Bits in",
        "Bits out",
        "Verdict",
        "Notes",
    ]
    table = ReportTable("Inactive Virtual Servers", headers, verdict_column=9)
    for path, verdict in sorted(analysis.virtual_verdicts.items()):
        if verdict.verdict == Verdict.IN_USE:
            continue
        virtual = parsed.virtuals[path]
        table.rows.append(
            [
                path,
                virtual.destination,
                virtual.default_pool,
                _join(virtual.irules),
                virtual.admin_state,
                virtual.availability,
                virtual.total_conns,
                virtual.bits_in,
                virtual.bits_out,
                verdict.verdict,
                verdict.notes,
            ]
        )
    return table


def _build_dead_chains(
    parsed: ParsedData, correlation: Correlation, analysis: AnalysisResult
) -> ReportTable:
    """One row per dead chain, grouped by pool: the artifact to take to the
    config owner. Commands are informational text, ordered VS -> pool ->
    nodes, and a node alive in another pool gets no delete line."""
    headers = [
        "Pool",
        "Partition",
        "Pool status",
        "Members",
        "Member statuses",
        "Node verdicts",
        "Virtual servers",
        "VS statuses",
        "VS verdicts",
        "Verdict",
        "Notes",
        "Suggested commands (informational)",
    ]
    table = ReportTable("Dead Chains", headers, verdict_column=9)
    for path, verdict in sorted(analysis.pool_verdicts.items()):
        if verdict.verdict != Verdict.OFFLINE_CANDIDATE:
            continue
        pool = parsed.pools[path]
        vs_paths = sorted(vs for vs, pools in correlation.virtual_to_pools.items() if path in pools)
        node_paths = sorted({member.node_full_path for member in pool.members})
        commands = []
        vs_states, vs_verdict_labels = [], []
        for vs_path in vs_paths:
            virtual = parsed.virtuals.get(vs_path)
            vs_states.append(
                f"{virtual.admin_state}/{virtual.availability or '?'}" if virtual else "?"
            )
            vs_verdict = analysis.virtual_verdicts.get(vs_path)
            label = vs_verdict.verdict if vs_verdict else ""
            vs_verdict_labels.append(label)
            if label == Verdict.OFFLINE_CANDIDATE:
                commands.append(f"tmsh delete ltm virtual {vs_path}")
        commands.append(f"tmsh delete ltm pool {path}")
        node_verdict_labels = []
        for node_path in node_paths:
            node_verdict = analysis.node_verdicts.get(node_path)
            label = node_verdict.verdict if node_verdict else ""
            node_verdict_labels.append(label)
            if label == Verdict.OFFLINE_CANDIDATE:
                commands.append(f"tmsh delete ltm node {node_path}")
        table.rows.append(
            [
                path,
                pool.partition,
                pool.availability,
                ", ".join(f"{m.node_full_path}:{m.port}" for m in pool.members),
                ", ".join(f"{m.admin_state}/{m.availability or '?'}" for m in pool.members),
                ", ".join(node_verdict_labels),
                ", ".join(vs_paths),
                ", ".join(vs_states),
                ", ".join(vs_verdict_labels),
                verdict.verdict,
                verdict.notes,
                "\n".join(commands),
            ]
        )
    return table


def _build_orphan_monitors(parsed: ParsedData, analysis: AnalysisResult) -> ReportTable:
    headers = [
        "Monitor",
        "Type",
        "Partition",
        "Verdict",
        "Notes",
        "Suggested command (informational)",
    ]
    table = ReportTable("Orphan Monitors", headers, verdict_column=3)
    for path, verdict in sorted(analysis.monitor_verdicts.items()):
        if verdict.verdict == Verdict.IN_USE:
            continue
        monitor = parsed.monitors[path]
        table.rows.append(
            [
                path,
                monitor.type,
                monitor.partition,
                verdict.verdict,
                verdict.notes,
                f"tmsh delete ltm monitor {monitor.type} {path}"
                if verdict.verdict == Verdict.ORPHAN
                else "",
            ]
        )
    return table


def _build_manual_review(analysis: AnalysisResult) -> ReportTable:
    headers = ["Object type", "Object", "Reason", "Caused by"]
    table = ReportTable("Manual Review", headers)
    for item in analysis.manual_review:
        table.rows.append([item.object_type, item.full_path, item.reason, item.caused_by])
    return table


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------


def write_xlsx(tables: dict[str, ReportTable], out_path: str) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)

    for table in tables.values():
        sheet = workbook.create_sheet(title=table.title[:31])
        sheet.append(table.headers)
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        for row in table.rows:
            sheet.append(row)
            if table.verdict_column is not None:
                verdict = row[table.verdict_column]
                color = VERDICT_FILLS.get(str(verdict))
                if color:
                    cell = sheet.cell(row=sheet.max_row, column=table.verdict_column + 1)
                    cell.fill = PatternFill("solid", fgColor=color)
        sheet.freeze_panes = "A2"
        last_column = get_column_letter(len(table.headers))
        sheet.auto_filter.ref = f"A1:{last_column}{max(sheet.max_row, 1)}"
        _autofit_columns(sheet, table)
    workbook.save(out_path)


def _autofit_columns(sheet, table: ReportTable) -> None:
    for index, header in enumerate(table.headers, start=1):
        width = len(str(header))
        for row in table.rows:
            value = row[index - 1]
            if value is not None:
                width = max(width, len(str(value)))
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, MAX_COLUMN_WIDTH)


def write_csv(tables: dict[str, ReportTable], out_dir: str) -> list[str]:
    directory = Path(out_dir)
    directory.mkdir(parents=True, exist_ok=True)
    written = []
    for key, table in tables.items():
        path = directory / f"{key}.csv"
        with open(path, "w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            writer.writerow(table.headers)
            writer.writerows(table.rows)
        written.append(str(path))
    return written
